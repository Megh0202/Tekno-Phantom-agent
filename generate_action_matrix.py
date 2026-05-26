"""
Generates action_matrix.xlsx — Tekno Phantom Agent capability tracker.
Run: python generate_action_matrix.py
"""

from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()
ws = wb.active
ws.title = "Action Matrix"

thin = Side(style="thin", color="CCCCCC")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

HEADERS    = ["Action", "Description", "VitaOne", "OrangeHRM", "AtoZBay", "Notes", "Last Tested"]
COL_WIDTHS = [30, 40, 14, 14, 14, 30, 14]

GROUP = "__GROUP__"

ROWS = [

    (GROUP, "NAVIGATION"),
    ("navigate",                  "Opens a URL in the browser"),
    ("navigate_back",             "Goes to the previous page"),
    ("navigate_forward",          "Goes to the next page"),
    ("refresh_page",              "Reloads the current page"),
    ("open_new_tab",              "Opens a URL in a new tab"),
    ("switch_tab",                "Switches to another open tab"),
    ("close_tab",                 "Closes the current tab"),
    ("switch_to_iframe",          "Enters an iframe to interact with it"),
    ("switch_to_main_frame",      "Exits iframe back to main page"),

    (GROUP, "MOUSE"),
    ("click",                     "Clicks any element on the page"),
    ("double_click",              "Double clicks an element"),
    ("right_click",               "Right clicks to open context menu"),
    ("middle_click",              "Middle clicks a link"),
    ("hover",                     "Moves mouse over an element"),
    ("drag_and_drop",             "Drags element and drops on target"),
    ("drag_by_offset",            "Drags element by pixel distance"),
    ("long_press",                "Holds click on element for a duration"),
    ("click_at_coordinates",      "Clicks at exact x y on the page"),

    (GROUP, "KEYBOARD"),
    ("key_press",                 "Presses a single key like Tab or Escape"),
    ("key_combination",           "Presses keys together like Ctrl+A"),
    ("press_enter",               "Presses the Enter key"),
    ("press_escape",              "Presses the Escape key"),
    ("press_tab",                 "Presses Tab to move to next field"),
    ("shortcut_select_all",       "Ctrl+A to select all"),
    ("shortcut_copy",             "Ctrl+C to copy"),
    ("shortcut_paste",            "Ctrl+V to paste"),
    ("shortcut_cut",              "Ctrl+X to cut"),
    ("shortcut_undo",             "Ctrl+Z to undo"),
    ("shortcut_redo",             "Ctrl+Y to redo"),

    (GROUP, "INPUT"),
    ("type",                      "Types text into an input field"),
    ("type_append",               "Adds text without clearing the field"),
    ("clear_input",               "Clears an input field"),
    ("focus_element",             "Focuses an element without typing"),
    ("blur_element",              "Removes focus from an element"),
    ("password_input",            "Types into a password field"),
    ("textarea_input",            "Types into a textarea"),
    ("type_in_contenteditable",   "Types in a rich text div field"),
    ("type_in_monaco_editor",     "Types inside Monaco code editor"),
    ("type_in_codemirror",        "Types inside CodeMirror editor"),

    (GROUP, "SELECTION"),
    ("select_option",             "Picks option from a native dropdown"),
    ("select_option_by_label",    "Picks dropdown option by visible text"),
    ("select_multi_option",       "Selects multiple options in a list"),
    ("deselect_option",           "Deselects an option in a list"),
    ("autocomplete_select",       "Types then picks from suggestions"),
    ("checkbox_check",            "Checks a checkbox"),
    ("checkbox_uncheck",          "Unchecks a checkbox"),
    ("checkbox_toggle",           "Toggles checkbox regardless of state"),
    ("radio_select",              "Selects a radio button"),
    ("toggle_switch_on",          "Turns a toggle switch on"),
    ("toggle_switch_off",         "Turns a toggle switch off"),

    (GROUP, "SCROLL"),
    ("scroll_page_down",          "Scrolls the page downward"),
    ("scroll_page_up",            "Scrolls the page upward"),
    ("scroll_horizontal",         "Scrolls left or right on the page"),
    ("scroll_element",            "Scrolls inside a container"),
    ("scroll_to_element",         "Scrolls until element is visible"),
    ("scroll_to_top",             "Jumps to the top of the page"),
    ("scroll_to_bottom",          "Jumps to the bottom of the page"),
    ("infinite_scroll_trigger",   "Scrolls to load more content"),

    (GROUP, "WAIT"),
    ("wait_timeout",              "Pauses for a fixed milliseconds"),
    ("wait_selector_visible",     "Waits until an element appears"),
    ("wait_selector_hidden",      "Waits until an element disappears"),
    ("wait_load_state",           "Waits for page load state"),
    ("wait_for_url_change",       "Waits until URL changes"),
    ("wait_for_text",             "Waits until element has a text"),
    ("wait_for_network_idle",     "Waits for all network calls to finish"),
    ("wait_for_download",         "Waits for a file download to begin"),

    (GROUP, "FILE"),
    ("upload_file",               "Uploads a file via file input"),
    ("upload_multiple_files",     "Uploads multiple files at once"),
    ("drag_file_upload",          "Drags a file onto an upload zone"),
    ("remove_uploaded_file",      "Removes an already uploaded file"),
    ("trigger_download",          "Clicks download and captures file"),
    ("verify_download",           "Asserts a file was downloaded"),

    (GROUP, "POPUP AND DIALOG"),
    ("handle_popup",              "Handles browser alert or confirm"),
    ("accept_dialog",             "Accepts a browser dialog"),
    ("dismiss_dialog",            "Dismisses a browser dialog"),
    ("close_modal",               "Closes a modal by clicking X"),
    ("confirm_modal",             "Clicks confirm inside a modal"),
    ("cancel_modal",              "Clicks cancel inside a modal"),
    ("wait_for_modal_open",       "Waits until a modal is visible"),
    ("wait_for_modal_close",      "Waits until a modal disappears"),

    (GROUP, "MENU"),
    ("open_menu",                 "Clicks to open a dropdown menu"),
    ("select_menu_item",          "Clicks an item inside a menu"),
    ("open_submenu",              "Hovers to reveal a nested submenu"),
    ("select_submenu_item",       "Clicks item inside a submenu"),
    ("open_context_menu",         "Right clicks to open context menu"),
    ("breadcrumb_navigate",       "Clicks a breadcrumb level"),

    (GROUP, "EXPAND AND COLLAPSE"),
    ("expand_accordion",          "Opens a collapsed section"),
    ("collapse_accordion",        "Closes an open section"),
    ("expand_sidebar",            "Opens the sidebar"),
    ("collapse_sidebar",          "Closes the sidebar"),
    ("expand_tree_node",          "Expands a node in a tree view"),
    ("collapse_tree_node",        "Collapses a node in a tree view"),

    (GROUP, "SPECIALIZED INPUTS"),
    ("set_date",                  "Picks a date in a date picker"),
    ("set_time",                  "Picks a time in a time picker"),
    ("set_datetime",              "Picks date and time together"),
    ("set_slider_value",          "Moves a slider to a value"),
    ("set_rating",                "Clicks a star rating widget"),
    ("pick_color",                "Selects a color from color picker"),
    ("increment_number",          "Clicks plus on a number spinner"),
    ("decrement_number",          "Clicks minus on a number spinner"),

    (GROUP, "TABLE AND GRID"),
    ("table_select_row",          "Clicks a table row to select it"),
    ("table_select_cell",         "Clicks a specific table cell"),
    ("table_edit_cell",           "Double clicks cell to edit inline"),
    ("table_sort_column",         "Clicks column header to sort"),
    ("table_filter",              "Applies filter on a table column"),
    ("table_paginate_next",       "Clicks next page in table"),
    ("table_paginate_prev",       "Clicks previous page in table"),
    ("table_go_to_page",          "Goes to a specific page number"),
    ("table_select_all_rows",     "Selects all rows in table"),
    ("table_deselect_all",        "Deselects all rows in table"),

    (GROUP, "VERIFICATION"),
    ("verify_text",               "Asserts text exists on the page"),
    ("verify_image",              "Compares screenshot to baseline"),
    ("verify_url",                "Asserts the current page URL"),
    ("verify_title",              "Asserts the page title"),
    ("verify_element_visible",    "Asserts element is visible"),
    ("verify_element_hidden",     "Asserts element is not visible"),
    ("verify_element_enabled",    "Asserts element is not disabled"),
    ("verify_input_value",        "Asserts current value of an input"),
    ("verify_checkbox_state",     "Asserts checkbox is checked or not"),
    ("verify_attribute",          "Asserts an HTML attribute value"),
    ("verify_toast_message",      "Asserts a toast notification showed"),
    ("verify_error_message",      "Asserts a validation error showed"),
    ("verify_element_count",      "Asserts number of elements found"),

    (GROUP, "READ AND CAPTURE"),
    ("read_text",                 "Reads element text into a variable"),
    ("read_input_value",          "Reads input value into a variable"),
    ("read_url",                  "Captures current URL"),
    ("capture_screenshot",        "Takes a screenshot of page or element"),

    (GROUP, "WINDOW"),
    ("resize_window",             "Sets the browser window size"),
    ("maximize_window",           "Maximizes the browser window"),
    ("set_viewport_mobile",       "Simulates mobile screen size"),
    ("set_viewport_tablet",       "Simulates tablet screen size"),

    (GROUP, "STORAGE AND AUTH"),
    ("set_cookie",                "Injects a cookie into the browser"),
    ("clear_cookies",             "Clears all browser cookies"),
    ("set_local_storage",         "Sets a localStorage key value"),
    ("clear_local_storage",       "Clears all localStorage data"),
    ("inject_auth_token",         "Sets auth token to skip login"),
    ("execute_script",            "Runs custom JavaScript on the page"),
    ("dispatch_event",            "Fires a DOM event on an element"),

    # ── CLICK ACTIONS (UI element specific) ───────────────────────────────────
    (GROUP, "CLICK ACTIONS"),
    ("click_button",              "Clicks a button element"),
    ("click_link",                "Clicks a hyperlink"),
    ("click_icon",                "Clicks an icon button"),
    ("click_image",               "Clicks an image element"),
    ("click_card",                "Clicks a card component"),
    ("click_list_item",           "Clicks an item in a list"),
    ("click_tab",                 "Clicks a UI tab to switch panel"),
    ("click_chip",                "Clicks a chip or tag element"),

    # ── FORM ACTIONS ──────────────────────────────────────────────────────────
    (GROUP, "FORM ACTIONS"),
    ("submit_form",               "Clicks the submit button of a form"),
    ("reset_form",                "Clicks reset to clear all form fields"),
    ("cancel_form",               "Clicks cancel and discards form changes"),
    ("stepper_next",              "Clicks Next in a multi-step wizard"),
    ("stepper_prev",              "Clicks Back in a multi-step wizard"),
    ("stepper_go_to_step",        "Clicks a specific step number in wizard"),
    ("otp_input",                 "Fills OTP boxes one by one"),
    ("phone_input",               "Types phone number with country code"),
    ("search_input",              "Types text in a search field"),
    ("clear_search",              "Clears the search input field"),
    ("copy_to_clipboard",         "Clicks copy button to copy a value"),

    # ── CHECKBOX, RADIO AND TOGGLE ────────────────────────────────────────────
    (GROUP, "CHECKBOX, RADIO AND TOGGLE"),
    ("checkbox_check_by_label",   "Checks a checkbox by its label text"),
    ("checkbox_uncheck_by_label", "Unchecks a checkbox by its label text"),
    ("radio_select_by_label",     "Selects a radio button by its label text"),
    ("toggle_switch_by_label",    "Turns a toggle on or off by its label"),
    ("verify_checkbox_checked",   "Asserts that a checkbox is checked"),
    ("verify_checkbox_unchecked", "Asserts that a checkbox is unchecked"),
    ("verify_radio_selected",     "Asserts a radio button is selected"),
    ("verify_toggle_on",          "Asserts a toggle switch is turned on"),
    ("verify_toggle_off",         "Asserts a toggle switch is turned off"),

    # ── TABS ──────────────────────────────────────────────────────────────────
    (GROUP, "TABS"),
    ("click_tab_by_label",        "Clicks a tab by its visible label text"),
    ("click_tab_by_index",        "Clicks a tab by its position number"),
    ("verify_active_tab",         "Asserts which tab is currently active"),

    # ── CHIPS AND TAGS ────────────────────────────────────────────────────────
    (GROUP, "CHIPS AND TAGS"),
    ("add_chip",                  "Adds a chip or tag to a tag input"),
    ("remove_chip",               "Removes a chip by clicking its X button"),
    ("add_tag",                   "Adds a tag label to an item"),
    ("remove_tag",                "Removes a tag label from an item"),
    ("select_chip_option",        "Clicks a chip option to select it"),
    ("deselect_chip_option",      "Clicks a selected chip to deselect it"),

    # ── CUSTOM DROPDOWN ───────────────────────────────────────────────────────
    (GROUP, "CUSTOM DROPDOWN"),
    ("open_custom_dropdown",      "Clicks to open a custom JS dropdown"),
    ("close_custom_dropdown",     "Closes an open custom dropdown"),
    ("select_custom_option",      "Picks an option from a custom dropdown"),
    ("combobox_search_select",    "Types in searchable dropdown then selects"),
    ("cascading_dropdown_select", "Selects from linked dependent dropdowns"),

    # ── EXPAND AND COLLAPSE (EXTENDED) ────────────────────────────────────────
    (GROUP, "EXPAND AND COLLAPSE EXTENDED"),
    ("expand_all_accordions",     "Expands all accordion sections at once"),
    ("collapse_all_accordions",   "Collapses all accordion sections"),
    ("expand_all_tree_nodes",     "Expands all nodes in a tree"),
    ("expand_row_details",        "Expands inline details row in a table"),
    ("collapse_row_details",      "Collapses inline details row in a table"),
    ("expand_details_section",    "Clicks to expand a details panel"),
    ("collapse_details_section",  "Clicks to collapse a details panel"),

    # ── NOTIFICATION AND ALERTS ───────────────────────────────────────────────
    (GROUP, "NOTIFICATION AND ALERTS"),
    ("dismiss_toast",             "Closes a toast notification"),
    ("dismiss_alert",             "Closes an alert or info banner"),
    ("dismiss_banner",            "Closes a top page banner message"),
    ("click_notification",        "Clicks a notification item"),
    ("mark_notification_read",    "Marks a notification as read"),
    ("clear_all_notifications",   "Clears all notifications at once"),

    # ── POPOVER, TOOLTIP AND DRAWER ───────────────────────────────────────────
    (GROUP, "POPOVER, TOOLTIP AND DRAWER"),
    ("open_popover",              "Clicks to open a popover"),
    ("close_popover",             "Closes an open popover"),
    ("open_tooltip",              "Hovers over element to show tooltip"),
    ("verify_tooltip_text",       "Asserts text shown in a tooltip"),
    ("open_drawer",               "Opens a slide-in drawer panel"),
    ("close_drawer",              "Closes a slide-in drawer panel"),

    # ── CAROUSEL AND SLIDER UI ────────────────────────────────────────────────
    (GROUP, "CAROUSEL AND SLIDER UI"),
    ("carousel_next",             "Clicks next on an image carousel"),
    ("carousel_prev",             "Clicks previous on an image carousel"),
    ("carousel_go_to_slide",      "Clicks a dot to jump to a specific slide"),

    # ── DATA AND VIEW ACTIONS ─────────────────────────────────────────────────
    (GROUP, "DATA AND VIEW ACTIONS"),
    ("apply_filter",              "Applies a filter from the filter panel"),
    ("clear_filter",              "Clears all applied filters"),
    ("sort_by",                   "Selects a sort order from sort dropdown"),
    ("switch_to_grid_view",       "Switches display to grid or card view"),
    ("switch_to_list_view",       "Switches display to list view"),
    ("switch_to_table_view",      "Switches display to table view"),
    ("bulk_select_items",         "Selects multiple items for bulk action"),
    ("apply_bulk_action",         "Applies an action to all selected items"),
    ("inline_edit_start",         "Clicks a field to start inline editing"),
    ("inline_edit_save",          "Saves an inline edit"),
    ("inline_edit_cancel",        "Cancels an inline edit"),
    ("pin_item",                  "Pins or favorites an item"),
    ("unpin_item",                "Unpins or unfavorites an item"),
    ("archive_item",              "Archives an item"),
    ("restore_item",              "Restores an archived item"),
    ("duplicate_item",            "Duplicates a record or item"),
    ("delete_item",               "Deletes an item after confirmation"),
    ("reorder_list_item",         "Drags an item to a new position in list"),
    ("resize_column",             "Drags column edge to resize it"),

    # ── TABLE EXTENDED ────────────────────────────────────────────────────────
    (GROUP, "TABLE EXTENDED"),
    ("table_deselect_row",        "Deselects a selected table row"),
    ("table_save_cell_edit",      "Saves inline cell edit with Enter"),
    ("table_cancel_cell_edit",    "Cancels inline cell edit with Escape"),
    ("table_clear_filter",        "Clears an applied column filter"),
    ("table_hide_column",         "Hides a table column"),
    ("table_show_column",         "Shows a hidden table column"),
    ("table_reorder_column",      "Drags column header to reorder"),
    ("table_group_by",            "Groups table rows by a column value"),
    ("table_bulk_action",         "Applies action to all selected rows"),
    ("table_export",              "Exports table data to file"),
    ("verify_table_row_count",    "Asserts the number of rows in a table"),

    # ── MEDIA ─────────────────────────────────────────────────────────────────
    (GROUP, "MEDIA"),
    ("media_play",                "Clicks play on a video or audio"),
    ("media_pause",               "Clicks pause on a video or audio"),
    ("media_seek",                "Seeks video to a specific timestamp"),
    ("media_mute",                "Mutes a video or audio element"),
    ("media_unmute",              "Unmutes a video or audio element"),
    ("media_fullscreen",          "Clicks fullscreen on a video player"),

    # ── VERIFICATION EXTENDED ─────────────────────────────────────────────────
    (GROUP, "VERIFICATION EXTENDED"),
    ("verify_tab_active",         "Asserts which tab panel is active"),
    ("verify_dropdown_value",     "Asserts selected value in a dropdown"),
    ("verify_badge_count",        "Asserts a badge or counter number"),
    ("verify_toggle_state",       "Asserts toggle is on or off"),
    ("verify_element_disabled",   "Asserts an element is disabled"),
    ("verify_placeholder",        "Asserts placeholder text of an input"),
    ("verify_list_item_count",    "Asserts number of items in a list"),
]

# ── Header row ────────────────────────────────────────────────────────────────
for ci, (h, w) in enumerate(zip(HEADERS, COL_WIDTHS), start=1):
    cell = ws.cell(row=1, column=ci, value=h)
    cell.fill = fill("333333")
    cell.font = Font(color="FFFFFF", bold=True, size=10)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = border
    ws.column_dimensions[get_column_letter(ci)].width = w
ws.row_dimensions[1].height = 22

# ── Data rows ─────────────────────────────────────────────────────────────────
row_num = 2
for entry in ROWS:

    if entry[0] == GROUP:
        _, label = entry
        ws.merge_cells(start_row=row_num, start_column=1,
                       end_row=row_num, end_column=len(HEADERS))
        cell = ws.cell(row=row_num, column=1, value=label)
        cell.fill = fill("EEEEEE")
        cell.font = Font(bold=True, size=10, color="333333")
        cell.alignment = Alignment(horizontal="left", vertical="center",
                                   indent=1)
        cell.border = border
        ws.row_dimensions[row_num].height = 16
        row_num += 1
        continue

    action, description = entry
    values = [action, description, "", "", "", "", ""]

    for ci, val in enumerate(values, start=1):
        cell = ws.cell(row=row_num, column=ci, value=val)
        cell.border = border
        cell.alignment = Alignment(horizontal="left", vertical="center",
                                   indent=1)
        cell.font = Font(size=10,
                         bold=(ci == 1),
                         color="000000")

    ws.row_dimensions[row_num].height = 15
    row_num += 1

ws.freeze_panes = "A2"

wb.save("action_matrix.xlsx")
print("Saved: action_matrix.xlsx")
